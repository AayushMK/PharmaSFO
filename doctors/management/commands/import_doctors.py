"""Import doctors from the MVTL MSL Universe spreadsheet.

Reads Sheet1 (the doctor "universe") and upserts Doctor rows keyed by NMC
number, creating Hospital + Area rows on the fly. Idempotent — safe to re-run.

Batched: instead of a get/update per row (thousands of round-trips, painfully
slow over a remote DB), it bulk-fetches existing Areas/Hospitals/Doctors and
uses bulk_create / bulk_update — a handful of queries total.

Usage (local, file mounted at /app):
    python manage.py import_doctors "MVTL MSL Universe.xlsx"
    python manage.py import_doctors "MVTL MSL Universe.xlsx" --dry-run
"""

from collections import Counter

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from doctors.models import Doctor, Hospital
from tour_plans.models import Area


def _clean(value):
    """Normalize a cell to a stripped string ('' for blanks; ints stay clean)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _split_hospitals(raw):
    """Split a '/'-separated hospital cell into de-duplicated, ordered tokens."""
    seen, tokens = set(), []
    for part in raw.split("/"):
        name = part.strip()[:255]
        if name and name.lower() not in seen:
            seen.add(name.lower())
            tokens.append(name)
    return tokens


DOCTOR_UPDATE_FIELDS = ["name", "hospital", "second_hospital", "area", "specialization", "phone"]


class Command(BaseCommand):
    help = "Import doctors from the MVTL MSL Universe .xlsx (Sheet1), in batches."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Path to the .xlsx file")
        parser.add_argument("--sheet", default="Sheet1", help="Worksheet name (default: Sheet1)")
        parser.add_argument("--dry-run", action="store_true", help="Roll back at the end; report only")

    def handle(self, *args, **opts):
        path, sheet, dry_run = opts["path"], opts["sheet"], opts["dry_run"]
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {path}")
        if sheet not in wb.sheetnames:
            raise CommandError(f"Sheet {sheet!r} not in {wb.sheetnames}")
        ws = wb[sheet]
        raw_rows = [r for r in ws.iter_rows(values_only=True) if any(_clean(c) for c in r)]
        wb.close()

        stats = Counter()
        # Column order: S.N. | Name | NMC | Speciality | City | Hospital | 2nd hosp/area | Phone
        parsed = {}  # nmc -> row dict (dedupe by NMC; last occurrence wins)
        for r in raw_rows[1:]:  # drop header
            cells = (list(r) + [None] * 8)[:8]
            _sn, name, nmc, spec, city, hosp_raw, second, phone = map(_clean, cells)
            stats["read"] += 1
            if not nmc:
                stats["skip_no_nmc"] += 1
                continue
            if not hosp_raw:
                stats["skip_no_hospital"] += 1
                continue
            if not city:
                stats["skip_no_city"] += 1
                continue
            tokens = _split_hospitals(hosp_raw)
            if not tokens:
                stats["skip_no_hospital"] += 1
                continue
            parsed[nmc] = {
                "nmc": nmc, "name": name, "spec": spec[:255], "city": city[:255],
                "tokens": tokens, "second": second[:255], "phone": phone[:20],
            }
        rows = list(parsed.values())

        with transaction.atomic():
            # 1) Areas — one per distinct city, matched case-insensitively so
            #    "Pokhara"/"POKHARA" collapse to one Area (as the row-by-row
            #    import did via a lower()-keyed cache).
            areas = {a.name.lower(): a for a in Area.objects.all()}
            canon = {}  # city_key(lower) -> canonical display name (first seen)
            for r in rows:
                r["city_key"] = r["city"].lower()
                canon.setdefault(r["city_key"], r["city"])
            new_areas = [Area(name=canon[k]) for k in canon if k not in areas]
            if new_areas:
                Area.objects.bulk_create(new_areas, ignore_conflicts=True)
                areas = {a.name.lower(): a for a in Area.objects.all()}
            stats["areas_created"] = len(new_areas)

            # 2) Hospitals — one per (token, doctor's-city area).
            needed = set()  # (name, area_id)
            for r in rows:
                area_id = areas[r["city_key"]].id
                for tok in r["tokens"]:
                    needed.add((tok, area_id))
            names = {n for n, _ in needed}
            hospitals = {
                (h.name, h.area_id): h
                for h in Hospital.objects.filter(name__in=names)
            }
            new_hospitals = [
                Hospital(name=n, area_id=a)
                for (n, a) in needed if (n, a) not in hospitals
            ]
            if new_hospitals:
                Hospital.objects.bulk_create(new_hospitals, ignore_conflicts=True)
                hospitals = {
                    (h.name, h.area_id): h
                    for h in Hospital.objects.filter(name__in=names)
                }
            stats["hospitals_created"] = len(new_hospitals)

            # 3) Doctors — split into create vs update, then two bulk statements.
            existing = {d.nmc_number: d for d in Doctor.objects.filter(nmc_number__in=parsed.keys())}
            to_create, to_update = [], []
            for r in rows:
                area_id = areas[r["city_key"]].id
                primary = hospitals[(r["tokens"][0], area_id)]
                fields = {
                    "name": r["name"], "hospital": primary, "second_hospital": r["second"],
                    "area": r["city"], "specialization": r["spec"], "phone": r["phone"],
                }
                if r["nmc"] in existing:
                    doc = existing[r["nmc"]]
                    for k, v in fields.items():
                        setattr(doc, k, v)
                    to_update.append(doc)
                else:
                    to_create.append(Doctor(nmc_number=r["nmc"], **fields))

            if to_create:
                Doctor.objects.bulk_create(to_create, batch_size=500)
            if to_update:
                Doctor.objects.bulk_update(to_update, DOCTOR_UPDATE_FIELDS, batch_size=200)

            if dry_run:
                transaction.set_rollback(True)

        self.stdout.write(self.style.MIGRATE_HEADING(
            "\nImport summary" + (" (DRY RUN — rolled back)" if dry_run else "")))
        self.stdout.write(f"  rows read (excl header):   {stats['read']}")
        self.stdout.write(f"  doctors created:           {len(to_create)}")
        self.stdout.write(f"  doctors updated:           {len(to_update)}")
        self.stdout.write(f"  skipped (no NMC):          {stats['skip_no_nmc']}")
        self.stdout.write(f"  skipped (no hospital):     {stats['skip_no_hospital']}")
        self.stdout.write(f"  skipped (no city):         {stats['skip_no_city']}")
        self.stdout.write(f"  hospitals created:         {stats['hospitals_created']}")
        self.stdout.write(f"  areas created:             {stats['areas_created']}")
        self.stdout.write(self.style.SUCCESS("  done."))
