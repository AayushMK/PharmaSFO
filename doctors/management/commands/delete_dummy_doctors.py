"""Delete leftover demo/dummy doctors — those NOT present in the MSL sheet.

The authoritative "real" doctor set is the import spreadsheet, keyed by NMC.
Any Doctor whose nmc_number is not in that set is treated as leftover demo
data and removed, together with its DailyCoverage rows (which are PROTECTed,
so they must go first) and — by cascade — its DoctorEmployeeRelation
assignments. Chemist/Stockist coverage is unrelated to doctors and untouched.

Run --dry-run first to see exactly what would be removed. Idempotent.

Usage (local, file mounted at /app):
    python manage.py delete_dummy_doctors "MVTL MSL Universe.xlsx" --dry-run
    python manage.py delete_dummy_doctors "MVTL MSL Universe.xlsx"
"""

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from daily_coverage.models import DailyCoverage
from doctor_employee_relation.models import DoctorEmployeeRelation
from doctors.models import Doctor


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


class Command(BaseCommand):
    help = "Delete demo doctors (and their coverage/assignments) not present in the MSL sheet."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Path to the MSL universe .xlsx (defines the real doctors)")
        parser.add_argument("--sheet", default="Sheet1", help="Worksheet name (default: Sheet1)")
        parser.add_argument("--dry-run", action="store_true", help="Report what would be deleted, then roll back")

    def handle(self, *args, **opts):
        path, sheet, dry_run = opts["path"], opts["sheet"], opts["dry_run"]
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"File not found: {path}")
        if sheet not in wb.sheetnames:
            raise CommandError(f"Sheet {sheet!r} not in {wb.sheetnames}")
        ws = wb[sheet]
        real_nmcs = set()
        for r in ws.iter_rows(values_only=True):
            nmc = _clean((list(r) + [None] * 8)[2])
            if nmc and nmc.lower() != "nmc no.":
                real_nmcs.add(nmc)
        wb.close()

        # Safety valve: if we parsed no NMCs, every doctor would look like a
        # dummy. Refuse rather than wipe the whole directory.
        if not real_nmcs:
            raise CommandError("No NMC numbers parsed from the sheet — refusing to run (would delete every doctor).")

        dummies = Doctor.objects.exclude(nmc_number__in=real_nmcs)

        with transaction.atomic():
            n_doctors = dummies.count()
            n_cov = DailyCoverage.objects.filter(doctor__in=dummies).count()
            n_rel = DoctorEmployeeRelation.objects.filter(doctor__in=dummies).count()
            sample = list(dummies.order_by("name").values_list("nmc_number", "name")[:80])

            # DailyCoverage.doctor is PROTECT — clear it before the doctors;
            # DoctorEmployeeRelation.doctor is CASCADE and goes automatically.
            DailyCoverage.objects.filter(doctor__in=dummies).delete()
            dummies.delete()

            if dry_run:
                transaction.set_rollback(True)

        header = "Would delete (DRY RUN — rolled back)" if dry_run else "Deleted"
        self.stdout.write(self.style.MIGRATE_HEADING(f"\n{header}:"))
        self.stdout.write(f"  real doctors in sheet (kept):   {len(real_nmcs)}")
        self.stdout.write(f"  dummy doctors:                  {n_doctors}")
        self.stdout.write(f"  their daily-coverage rows:      {n_cov}")
        self.stdout.write(f"  their assignments (cascade):    {n_rel}")
        if n_doctors:
            self.stdout.write("\n  doctors removed:")
            for nmc, name in sample:
                self.stdout.write(f"    [{nmc or 'no-nmc'}] Dr. {name}")
            if n_doctors > len(sample):
                self.stdout.write(f"    … and {n_doctors - len(sample)} more")
        self.stdout.write(self.style.SUCCESS("\n  done."))
